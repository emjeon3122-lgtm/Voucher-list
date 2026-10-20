"""운영: 설정 검사, Graph 메일 발송, 자동 백업, 명단 엑셀."""
import json
import os
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app import backup, create_app, mailer
from app.config import check_production
from conftest import login


def test_production_config_refuses_weak_settings(tmp_path):
    with pytest.raises(RuntimeError) as e:
        create_app({"DATA_DIR": str(tmp_path), "SECRET_KEY": "short", "AUTH_MODE": "oidc",
                    "OIDC_CLIENT_ID": "", "PUBLIC_URL": "http://x", "ALLOWED_EMAIL_DOMAINS": []})
    msg = str(e.value)
    assert "SESSION_SECRET" in msg and "OIDC_TENANT_ID" in msg and "https" in msg and "ALLOWED_EMAIL_DOMAINS" in msg


def test_production_config_ok():
    check_production({
        "SECRET_KEY": "x" * 48, "AUTH_MODE": "oidc", "OIDC_TENANT_ID": "t", "OIDC_ISSUER": "", "OIDC_CLIENT_ID": "c",
        "OIDC_CLIENT_SECRET": "s", "PUBLIC_URL": "https://192.168.100.25:7031", "ALLOWED_EMAIL_DOMAINS": ["bdo.kr"],
        "MAIL_BACKEND": "graph", "MAIL_SENDER": "voucher@bdo.kr", "GRAPH_TENANT_ID": "t", "GRAPH_CLIENT_ID": "c",
        "GRAPH_CLIENT_SECRET": "s", "SMTP_HOST": "",
    })


class FakeResp:
    def __init__(self, data=b"{}"):
        self.data = data

    def read(self):
        return self.data

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def test_graph_mail_backend(app, monkeypatch):
    calls = []

    def fake_urlopen(req, timeout=None):
        calls.append(req)
        if "oauth2/v2.0/token" in req.full_url:
            return FakeResp(json.dumps({"access_token": "tok", "expires_in": 3600}).encode())
        return FakeResp()

    monkeypatch.setattr(mailer.urllib.request, "urlopen", fake_urlopen)
    monkeypatch.setitem(mailer._token, "value", None)
    app.config.update(MAIL_BACKEND="graph", MAIL_SENDER="voucher@co.kr", GRAPH_TENANT_ID="t",
                      GRAPH_CLIENT_ID="c", GRAPH_CLIENT_SECRET="s")
    with app.test_request_context():
        msg = mailer.build("kim@co.kr", "제목", "email/login_code", user={"name": "김철수"}, code="123456", minutes=10)
        mailer.send_now([msg, msg])
    token_calls = [c for c in calls if "token" in c.full_url]
    send_calls = [c for c in calls if "sendMail" in c.full_url]
    assert len(token_calls) == 1 and len(send_calls) == 2  # 토큰은 재사용
    assert send_calls[0].full_url == "https://graph.microsoft.com/v1.0/users/voucher%40co.kr/sendMail"
    assert send_calls[0].get_header("Authorization") == "Bearer tok"
    body = json.loads(send_calls[0].data)
    assert body["message"]["toRecipients"][0]["emailAddress"]["address"] == "kim@co.kr"
    assert "123456" in body["message"]["body"]["content"]


def test_daily_backup_and_prune(app):
    cfg = app.config
    d = backup.backup_dir(cfg)
    os.makedirs(d, exist_ok=True)
    old = os.path.join(d, "voucher-backup-20250101.db")
    open(old, "w").close()
    assert backup.run_daily(cfg, datetime(2026, 10, 20, 2, 0)) is None  # BACKUP_HOUR(3시) 전
    made = backup.run_daily(cfg, datetime(2026, 10, 20, 3, 5))
    assert made and made.endswith("voucher-backup-20261020.db") and os.path.getsize(made) > 0
    assert backup.run_daily(cfg, datetime(2026, 10, 20, 9, 0)) is None  # 하루 한 번
    assert not os.path.exists(old)  # 보관 기간 지난 사본 삭제


def test_roster_excel_roundtrip(app, outbox):
    c = app.test_client()
    login(c, outbox, "admin@co.kr")
    wb = load_workbook(BytesIO(c.get("/admin/users/export.xlsx").data))
    assert wb.active["A1"].value == "이메일" and wb.active.max_row == 4

    up = Workbook()
    ws = up.active
    ws.append(["이메일", "성명", "부서", "사번", "권한"])
    ws.append(["kim@co.kr", "김철수", "전략팀", "1001", ""])
    ws.append(["new@co.kr", "신입", "총무팀", "2001", "관리자"])
    ws.append(["bad-email", "오류", "", "", ""])
    bio = BytesIO()
    up.save(bio)
    bio.seek(0)
    c.post("/admin/users/import", data={"file": (bio, "명단.xlsx")}, content_type="multipart/form-data")
    from app.db import connect, db_path
    conn = connect(db_path(app))
    got = {r["email"]: (r["department"], r["role"]) for r in conn.execute("SELECT * FROM users")}
    conn.close()
    assert got["kim@co.kr"] == ("전략팀", "staff")
    assert got["new@co.kr"] == ("총무팀", "admin")
    assert "lee@co.kr" in got  # 엑셀에 없는 직원은 지우지 않음
