from io import BytesIO

from openpyxl import load_workbook

from app.db import connect, db_path
from conftest import login, sign_link, signature_data_url

T50, T10 = 1, 2  # voucher_types id
KIM, LEE = 2, 3  # users id


def q(app, sql, *args):
    conn = connect(db_path(app))
    try:
        return conn.execute(sql, args).fetchall()
    finally:
        conn.close()


def add_in(c, qty, type_id=T50, date="2026-01-05"):
    return c.post("/ledger/in", data={"txn_date": date, "voucher_type_id": type_id, "quantity": qty})


def add_out(c, qty, recipient=KIM, type_id=T50, date="2026-01-10", send=False):
    data = {"txn_date": date, "voucher_type_id": type_id, "quantity": qty, "recipient_id": recipient,
            "purpose": "포상"}
    if send:
        data["send_request"] = "1"
    return c.post("/ledger/out", data=data)


def test_login_requires_registered_email(app, outbox):
    c = app.test_client()
    resp = c.post("/login", data={"email": "nobody@co.kr"})
    assert resp.status_code == 400
    assert not outbox
    assert c.get("/ledger").status_code == 302  # 로그인 필요


def test_wrong_code_rejected(app, outbox):
    c = app.test_client()
    c.post("/login", data={"email": "kim@co.kr"})
    assert c.post("/login/verify", data={"code": "000000x"}).status_code in (400, 302)
    assert c.get("/my").status_code == 302


def test_staff_cannot_open_ledger(app, outbox):
    c = app.test_client()
    login(c, outbox, "kim@co.kr")
    assert c.get("/ledger").status_code == 403
    assert c.get("/my").status_code == 200


def test_stock_and_balance(app, admin_client):
    add_in(admin_client, 10)
    add_in(admin_client, 5, type_id=T10)
    add_out(admin_client, 3)
    rows = q(app, "SELECT direction, quantity, amount FROM transactions ORDER BY id")
    assert [tuple(r) for r in rows] == [("IN", 10, 500000), ("IN", 5, 50000), ("OUT", 3, 150000)]
    html = admin_client.get("/ledger").get_data(as_text=True)
    assert "400,000" in html  # 총 잔액 = 500,000 + 50,000 - 150,000
    assert "7<small>매</small>" in html  # 5만원권 재고


def test_out_cannot_exceed_stock(app, admin_client):
    add_in(admin_client, 2)
    add_out(admin_client, 3)
    assert q(app, "SELECT COUNT(*) FROM transactions WHERE direction='OUT'")[0][0] == 0


def test_backdated_out_cannot_make_history_negative(app, admin_client):
    add_in(admin_client, 5, date="2026-01-05")
    add_out(admin_client, 5, date="2026-01-10")
    add_in(admin_client, 5, date="2026-01-20")
    # 1/15 기준 재고는 0 → 출고 불가 (현재 총재고 5 이지만 과거 잔액이 음수가 됨)
    add_out(admin_client, 1, date="2026-01-15")
    assert q(app, "SELECT COUNT(*) FROM transactions WHERE direction='OUT'")[0][0] == 1


def test_void_in_blocked_when_stock_used(app, admin_client):
    add_in(admin_client, 5)
    add_out(admin_client, 4)
    admin_client.post("/ledger/1/void", data={"reason": "오입력"})
    assert q(app, "SELECT is_void FROM transactions WHERE id=1")[0][0] == 0
    admin_client.post("/ledger/2/void", data={"reason": "오입력"})
    admin_client.post("/ledger/1/void", data={"reason": "오입력"})
    assert [r[0] for r in q(app, "SELECT is_void FROM transactions ORDER BY id")] == [1, 1]


def test_sign_flow(app, admin_client, outbox):
    add_in(admin_client, 10)
    add_out(admin_client, 2, send=True)
    msg = outbox[-1]
    assert msg["To"] == "kim@co.kr"
    link = sign_link(msg)
    assert q(app, "SELECT sign_status FROM transactions WHERE id=2")[0][0] == "REQUESTED"
    assert q(app, "SELECT mail_sent_at IS NOT NULL FROM signature_requests")[0][0] == 1

    anon = app.test_client()
    page = anon.get(link)
    assert page.status_code == 200 and "김철수" in page.get_data(as_text=True)

    # 이름 불일치 / 빈 서명 거부
    r = anon.post(link, data={"agree": "1", "signer_name": "홍길동", "signature": signature_data_url()})
    assert r.status_code == 400
    r = anon.post(link, data={"agree": "1", "signer_name": "김철수", "signature": signature_data_url(blank=True)})
    assert r.status_code == 400
    assert q(app, "SELECT sign_status FROM transactions WHERE id=2")[0][0] == "REQUESTED"

    r = anon.post(link, data={"agree": "1", "signer_name": "김 철수", "signature": signature_data_url()})
    assert r.status_code == 302
    assert q(app, "SELECT sign_status FROM transactions WHERE id=2")[0][0] == "SIGNED"
    req = q(app, "SELECT status, signature_file, signature_sha256 FROM signature_requests")[0]
    assert req["status"] == "SIGNED" and req["signature_file"] and len(req["signature_sha256"]) == 64
    assert "서명이 완료되었습니다" in anon.get(link).get_data(as_text=True)

    img = admin_client.get("/ledger").get_data(as_text=True)
    assert "/signature/2.png" in img
    assert admin_client.get("/signature/2.png").mimetype == "image/png"

    other = app.test_client()
    login(other, outbox, "lee@co.kr")
    assert other.get("/signature/2.png").status_code == 403


def test_bulk_request_groups_by_recipient_and_resend_invalidates_old_link(app, admin_client, outbox):
    add_in(admin_client, 10)
    add_out(admin_client, 1, recipient=KIM)
    add_out(admin_client, 1, recipient=KIM)
    add_out(admin_client, 1, recipient=LEE)
    n = len(outbox)
    admin_client.post("/ledger/sign-request", data={"txn_ids": ["2", "3", "4"]})
    assert len(outbox) == n + 2
    old = sign_link([m for m in outbox if m["To"] == "kim@co.kr"][-1])

    admin_client.post("/ledger/sign-request", data={"txn_ids": ["2", "3"]})
    new = sign_link(outbox[-1])
    c = app.test_client()
    assert c.get(old).status_code == 410
    assert c.get(new).status_code == 200
    c.post(new, data={"agree": "1", "signer_name": "김철수", "signature": signature_data_url()})
    statuses = [r[0] for r in q(app, "SELECT sign_status FROM transactions WHERE id IN (2,3,4) ORDER BY id")]
    assert statuses == ["SIGNED", "SIGNED", "REQUESTED"]


def test_staff_can_sign_from_my_page(app, admin_client, outbox):
    add_in(admin_client, 3)
    add_out(admin_client, 1, recipient=LEE, send=True)
    c = app.test_client()
    login(c, outbox, "lee@co.kr")
    assert "지금 서명하기" in c.get("/my").get_data(as_text=True)
    r = c.post("/sign/r/1", data={"agree": "1", "signer_name": "이영희", "signature": signature_data_url()})
    assert r.status_code == 302
    assert q(app, "SELECT sign_status FROM transactions WHERE id=2")[0][0] == "SIGNED"
    kim = app.test_client()
    login(kim, outbox, "kim@co.kr")
    assert kim.get("/sign/r/1").status_code == 404


def test_excel_export(app, admin_client, outbox):
    add_in(admin_client, 4)
    add_out(admin_client, 1, send=True)
    app.test_client().post(sign_link(outbox[-1]),
                           data={"agree": "1", "signer_name": "김철수", "signature": signature_data_url()})
    resp = admin_client.get("/ledger/export.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    ws = wb["수불대장"]
    assert ws["C3"].value == "출고" and ws["J3"].value == 150000
    assert len(ws._images) == 1
    assert wb["재고현황"]["E2"].value == 3


def test_csrf_enforced(tmp_path):
    from app import create_app
    app = create_app({"DATA_DIR": str(tmp_path), "MAIL_BACKEND": "memory", "MAIL_ASYNC": False})
    c = app.test_client()
    assert c.post("/login", data={"email": "a@b.co"}).status_code == 400
    assert c.get("/healthz").data == b"ok"
