"""수불 대장 업무 로직: 입고/출고, 재고·잔액 계산, 취소, 전자서명 요청/완료, 엑셀 내보내기."""
import base64
import os
import re
import secrets
from datetime import timedelta
from io import BytesIO

from flask import current_app, g
from PIL import Image, ImageOps

from .db import write_tx
from .utils import client_ip, now, now_str, parse_int, sha256_hex, today_str, valid_date

MAX_QTY = 100_000


class LedgerError(Exception):
    """화면에 그대로 보여줄 업무 오류."""


def audit(db, action, target="", detail="", user_id=None):
    if user_id is None and getattr(g, "user", None) is not None:
        user_id = g.user["id"]
    db.execute(
        "INSERT INTO audit_log (user_id, action, target, detail, ip, created_at) VALUES (?, ?, ?, ?, ?, ?)",
        (user_id, action, target, detail[:1000], client_ip(), now_str()),
    )


def signature_dir():
    path = os.path.join(current_app.config["DATA_DIR"], "signatures")
    os.makedirs(path, exist_ok=True)
    return path


# --------------------------------------------------------------------------- 재고 / 대장 조회

def stock_summary(db):
    rows = db.execute(
        """
        SELECT vt.id, vt.name, vt.face_value, vt.is_active,
               COALESCE(SUM(CASE WHEN t.direction = 'IN'  THEN t.quantity END), 0) AS in_qty,
               COALESCE(SUM(CASE WHEN t.direction = 'OUT' THEN t.quantity END), 0) AS out_qty
          FROM voucher_types vt
          LEFT JOIN transactions t ON t.voucher_type_id = vt.id AND t.is_void = 0
         GROUP BY vt.id
         ORDER BY vt.sort_order, vt.name, vt.face_value DESC
        """
    ).fetchall()
    result = []
    for r in rows:
        qty = r["in_qty"] - r["out_qty"]
        result.append(dict(r, qty=qty, balance=qty * r["face_value"]))
    return result


def _type_running(db, type_id):
    """해당 권종의 거래를 (일자, 등록순) 으로 정렬한 누적 재고."""
    return db.execute(
        """
        SELECT id, txn_date,
               SUM(CASE direction WHEN 'IN' THEN quantity ELSE -quantity END)
                   OVER (ORDER BY txn_date, id) AS bal
          FROM transactions
         WHERE voucher_type_id = ? AND is_void = 0
         ORDER BY txn_date, id
        """,
        (type_id,),
    ).fetchall()


def available_qty(db, type_id, txn_date):
    """txn_date 일자로 새 출고를 넣을 때 출고 가능한 최대 수량.
    과거 일자로 출고하더라도 그 이후 어느 시점의 재고도 음수가 되지 않도록 검사합니다."""
    bal_at, min_after = 0, None
    for r in _type_running(db, type_id):
        if r["txn_date"] <= txn_date:
            bal_at = r["bal"]
        else:
            min_after = r["bal"] if min_after is None else min(min_after, r["bal"])
    return bal_at if min_after is None else min(bal_at, min_after)


LEDGER_SQL = """
WITH base AS (
    SELECT t.*,
           SUM(CASE WHEN t.is_void = 1 THEN 0 WHEN t.direction = 'IN' THEN t.quantity ELSE -t.quantity END)
               OVER (PARTITION BY t.voucher_type_id ORDER BY t.txn_date, t.id) AS qty_balance,
           SUM(CASE WHEN t.is_void = 1 THEN 0 WHEN t.direction = 'IN' THEN t.amount ELSE -t.amount END)
               OVER (ORDER BY t.txn_date, t.id) AS amount_balance
      FROM transactions t
)
SELECT b.*,
       vt.name AS type_name,
       gv.name AS giver_name,
       COALESCE(b.rcpt_name, rc.name) AS recipient_name, COALESCE(b.rcpt_email, rc.email) AS recipient_email,
       COALESCE(b.rcpt_dept, rc.department) AS recipient_dept, rc.job_title AS recipient_title,
       sr.status AS req_status, sr.requested_at, sr.expires_at, sr.mail_sent_at, sr.mail_error,
       sr.signed_at, sr.signer_ip, sr.signature_file, sr.signature_sha256
  FROM base b
  JOIN voucher_types vt ON vt.id = b.voucher_type_id
  LEFT JOIN users gv ON gv.id = b.giver_id
  LEFT JOIN users rc ON rc.id = b.recipient_id
  LEFT JOIN signature_requests sr ON sr.id = b.sign_request_id
"""


def ledger_rows(db, f):
    """대장 행. 누적 재고/잔액은 전체 기준으로 계산한 뒤 필터를 적용합니다."""
    where, params = [], []
    if valid_date(f.get("date_from")):
        where.append("b.txn_date >= ?")
        params.append(f["date_from"])
    if valid_date(f.get("date_to")):
        where.append("b.txn_date <= ?")
        params.append(f["date_to"])
    if parse_int(f.get("type_id")):
        where.append("b.voucher_type_id = ?")
        params.append(parse_int(f["type_id"]))
    if f.get("direction") in ("IN", "OUT"):
        where.append("b.direction = ?")
        params.append(f["direction"])
    if f.get("sign") in ("UNREQUESTED", "REQUESTED", "SIGNED"):
        where.append("b.sign_status = ?")
        params.append(f["sign"])
    if f.get("recipient_id"):
        where.append("b.recipient_id = ?")
        params.append(f["recipient_id"])
    if f.get("hide_void"):
        where.append("b.is_void = 0")
    if f.get("q"):
        like = f"%{f['q']}%"
        where.append(
            "(b.rcpt_name LIKE ? OR b.rcpt_dept LIKE ? OR b.rcpt_email LIKE ? OR gv.name LIKE ? "
            "OR b.purpose LIKE ? OR b.serial_no LIKE ? OR b.counterparty LIKE ? OR vt.name LIKE ?)"
        )
        params.extend([like] * 8)
    sql = LEDGER_SQL + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY b.txn_date, b.id"
    return db.execute(sql, params).fetchall()


def ledger_totals(rows):
    t = dict(in_qty=0, in_amount=0, out_qty=0, out_amount=0)
    for r in rows:
        if r["is_void"]:
            continue
        key = "in" if r["direction"] == "IN" else "out"
        t[key + "_qty"] += r["quantity"]
        t[key + "_amount"] += r["amount"]
    return t


def ledger_version(db):
    """대장 변경 감지용 값 (화면 자동 새로고침)."""
    r = db.execute("SELECT COUNT(*) AS n, COALESCE(MAX(updated_at), '') AS u FROM transactions").fetchone()
    return f"{r['n']}|{r['u']}"


# --------------------------------------------------------------------------- 입고 / 출고 / 취소

def _parse_common(db, form):
    txn_date = (form.get("txn_date") or "").strip()
    if not valid_date(txn_date):
        raise LedgerError("일자를 올바르게 입력하세요. (YYYY-MM-DD)")
    if txn_date > today_str():
        raise LedgerError("미래 일자로는 등록할 수 없습니다.")
    vt = db.execute(
        "SELECT * FROM voucher_types WHERE id = ? AND is_active = 1", (parse_int(form.get("voucher_type_id")),)
    ).fetchone()
    if vt is None:
        raise LedgerError("상품권 종류/권종을 선택하세요.")
    qty = parse_int(form.get("quantity"))
    if qty is None or not 0 < qty <= MAX_QTY:
        raise LedgerError("수량은 1 이상의 숫자로 입력하세요.")
    return dict(
        txn_date=txn_date,
        vt=vt,
        quantity=qty,
        serial_no=(form.get("serial_no") or "").strip()[:500],
        purpose=(form.get("purpose") or "").strip()[:500],
    )


def upsert_person(db, person):
    """회사 주소록에서 고른 사람을 저장(또는 최신 정보로 갱신)하고 id 를 돌려줍니다. 쓰기 트랜잭션 안에서 호출."""
    ts = now_str()
    row = (db.execute("SELECT id FROM users WHERE directory_id = ?", (person["id"],)).fetchone()
           or db.execute("SELECT id FROM users WHERE email = ? AND directory_id IS NULL", (person["email"],)).fetchone())
    values = (person["id"], person["email"], person["name"], person["department"], person["job_title"],
              person["employee_no"], ts)
    if row:
        # 메일이 다른 사람에게 넘어간 경우(퇴사자 메일 재사용 등) 옛 기록의 메일을 비워 충돌을 피함
        db.execute("UPDATE users SET email = email || '#old' || id WHERE email = ? AND id != ?", (person["email"], row["id"]))
        db.execute("""UPDATE users SET directory_id=?, email=?, name=?, department=?, job_title=?, employee_no=?,
                             is_active=1, updated_at=? WHERE id=?""", (*values, row["id"]))
        return row["id"]
    db.execute("UPDATE users SET email = email || '#old' || id WHERE email = ?", (person["email"],))
    return db.execute(
        """INSERT INTO users (directory_id, email, name, department, job_title, employee_no, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", (*values, ts),
    ).lastrowid


def refresh_people(db, people):
    """주소록에서 다시 확인한 결과를 반영합니다. people: {users.id: 주소록 정보 또는 None(퇴사·사용 중지)}"""
    with write_tx(db):
        for user_id, person in people.items():
            if person is None:
                db.execute("UPDATE users SET is_active=0, updated_at=? WHERE id=?", (now_str(), user_id))
            else:
                upsert_person(db, person)


def create_transaction(db, direction, form, actor, recipient=None):
    """recipient: 출고일 때 회사 주소록에서 확인한 수령자 정보 (directory.get 결과)."""
    data = _parse_common(db, form)
    vt, ts = data["vt"], now_str()
    giver_id = recipient_id = None
    counterparty = ""
    if direction == "IN":
        counterparty = (form.get("counterparty") or "").strip()[:200]
        giver_id = actor["id"]
        sign_status = "NONE"
    else:
        if not recipient:
            raise LedgerError("수령자를 회사 주소록에서 검색해 선택하세요.")
        giver_id = actor["id"]
        sign_status = "UNREQUESTED"

    with write_tx(db):
        if direction == "OUT":
            recipient_id = upsert_person(db, recipient)
            avail = available_qty(db, vt["id"], data["txn_date"])
            if data["quantity"] > avail:
                raise LedgerError(
                    f"재고가 부족합니다. {vt['name']} {vt['face_value']:,}원권 "
                    f"{data['txn_date']} 기준 출고 가능 수량: {max(avail, 0):,}매"
                )
        cur = db.execute(
            """
            INSERT INTO transactions (txn_date, direction, voucher_type_id, quantity, face_value, amount,
                                      serial_no, counterparty, giver_id, recipient_id, rcpt_name, rcpt_email, rcpt_dept,
                                      purpose, sign_status, created_by, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (data["txn_date"], direction, vt["id"], data["quantity"], vt["face_value"],
             data["quantity"] * vt["face_value"], data["serial_no"], counterparty, giver_id, recipient_id,
             recipient and recipient["name"], recipient and recipient["email"], recipient and recipient["department"],
             data["purpose"], sign_status, actor["id"], ts, ts),
        )
        txn_id = cur.lastrowid
        to = f" → {recipient['name']} <{recipient['email']}>" if recipient else ""
        audit(db, "TXN_" + direction, f"txn:{txn_id}",
              f"{data['txn_date']} {vt['name']} {vt['face_value']:,}원 x {data['quantity']}{to}")
    return txn_id


def void_transaction(db, txn_id, reason, actor):
    reason = (reason or "").strip()
    if not reason:
        raise LedgerError("취소 사유를 입력하세요.")
    with write_tx(db):
        t = db.execute("SELECT * FROM transactions WHERE id = ?", (txn_id,)).fetchone()
        if t is None or t["is_void"]:
            raise LedgerError("이미 취소되었거나 존재하지 않는 내역입니다.")
        if t["direction"] == "IN":
            rows = _type_running(db, t["voucher_type_id"])
            idx = next(i for i, r in enumerate(rows) if r["id"] == t["id"])
            if min(r["bal"] for r in rows[idx:]) - t["quantity"] < 0:
                raise LedgerError("이 입고를 취소하면 이후 재고가 음수가 됩니다. 관련 출고를 먼저 취소하세요.")
        ts = now_str()
        db.execute(
            "UPDATE transactions SET is_void=1, void_reason=?, voided_by=?, voided_at=?, updated_at=? WHERE id=?",
            (reason[:500], actor["id"], ts, ts, txn_id),
        )
        # 대기 중인 서명 요청에 남은 유효 항목이 없으면 요청도 취소
        if t["sign_status"] == "REQUESTED" and t["sign_request_id"]:
            left = db.execute(
                """SELECT COUNT(*) FROM signature_request_items i JOIN transactions x ON x.id = i.transaction_id
                    WHERE i.request_id = ? AND x.is_void = 0""",
                (t["sign_request_id"],),
            ).fetchone()[0]
            if left == 0:
                db.execute("UPDATE signature_requests SET status='CANCELLED' WHERE id=? AND status='PENDING'",
                           (t["sign_request_id"],))
        audit(db, "TXN_VOID", f"txn:{txn_id}", reason)


# --------------------------------------------------------------------------- 전자서명

def create_sign_requests(db, txn_ids, actor):
    """선택한 출고 건에 대해 수령자별로 서명 요청을 1건씩 만듭니다. 메일 발송용 정보 목록을 반환."""
    ids = sorted({i for i in (parse_int(x) for x in txn_ids) if i})
    if not ids:
        raise LedgerError("서명 요청할 출고 내역을 선택하세요.")
    marks = ",".join("?" * len(ids))
    results = []
    with write_tx(db):
        rows = db.execute(
            f"""SELECT t.*, vt.name AS type_name, u.email AS recipient_email, u.name AS recipient_name,
                       u.department AS recipient_dept, u.is_active AS recipient_active
                  FROM transactions t
                  JOIN voucher_types vt ON vt.id = t.voucher_type_id
                  LEFT JOIN users u ON u.id = t.recipient_id
                 WHERE t.id IN ({marks}) ORDER BY t.txn_date, t.id""",
            ids,
        ).fetchall()
        if len(rows) != len(ids):
            raise LedgerError("존재하지 않는 내역이 포함되어 있습니다.")
        for r in rows:
            label = f"No.{r['id']}"
            if r["direction"] != "OUT" or r["is_void"]:
                raise LedgerError(f"{label}: 유효한 출고 내역만 서명 요청할 수 있습니다.")
            if r["sign_status"] == "SIGNED":
                raise LedgerError(f"{label}: 이미 서명이 완료된 내역입니다.")
            if not r["recipient_active"]:
                raise LedgerError(f"{label}: {r['recipient_name']}님은 회사 주소록에서 퇴사(계정 사용 중지)로 확인되어 "
                                  f"메일을 보낼 수 없습니다.")

        groups = {}
        for r in rows:
            groups.setdefault(r["recipient_id"], []).append(r)

        ts = now()
        expires = now_str(ts + timedelta(days=current_app.config["SIGN_LINK_DAYS"]))
        for recipient_id, items in groups.items():
            item_ids = [r["id"] for r in items]
            item_marks = ",".join("?" * len(item_ids))
            # 재발송: 기존 대기 요청은 취소하고(옛 링크 무효화), 그 요청에만 있던 다른 건은 '미요청'으로 되돌림
            for old_id in {r["sign_request_id"] for r in items if r["sign_status"] == "REQUESTED"}:
                db.execute("UPDATE signature_requests SET status='CANCELLED' WHERE id=? AND status='PENDING'",
                           (old_id,))
                db.execute(
                    f"""UPDATE transactions SET sign_status='UNREQUESTED', sign_request_id=NULL, updated_at=?
                         WHERE sign_request_id=? AND sign_status='REQUESTED' AND id NOT IN ({item_marks})""",
                    [now_str(ts), old_id, *item_ids],
                )
            token = secrets.token_urlsafe(32)
            req_id = db.execute(
                """INSERT INTO signature_requests (token_hash, recipient_id, email, status, expires_at,
                                                  requested_by, requested_at)
                   VALUES (?, ?, ?, 'PENDING', ?, ?, ?)""",
                (sha256_hex(token), recipient_id, items[0]["recipient_email"], expires, actor["id"], now_str(ts)),
            ).lastrowid
            db.executemany("INSERT INTO signature_request_items (request_id, transaction_id) VALUES (?, ?)",
                           [(req_id, i) for i in item_ids])
            db.execute(
                f"""UPDATE transactions SET sign_status='REQUESTED', sign_request_id=?, updated_at=?
                     WHERE id IN ({item_marks})""",
                [req_id, now_str(ts), *item_ids],
            )
            audit(db, "SIGN_REQUEST", f"sigreq:{req_id}",
                  f"{items[0]['recipient_name']} <{items[0]['recipient_email']}> / 내역 {item_ids}")
            results.append(dict(
                request_id=req_id, token=token, email=items[0]["recipient_email"],
                name=items[0]["recipient_name"], department=items[0]["recipient_dept"],
                lines=items, total=sum(r["amount"] for r in items), expires_at=expires,
            ))
    return results


def find_request_by_token(db, token):
    if not token or len(token) > 200:
        return None
    return db.execute("SELECT * FROM signature_requests WHERE token_hash = ?", (sha256_hex(token),)).fetchone()


def request_lines(db, req_id):
    return db.execute(
        """SELECT t.*, vt.name AS type_name, gv.name AS giver_name
             FROM signature_request_items i
             JOIN transactions t ON t.id = i.transaction_id
             JOIN voucher_types vt ON vt.id = t.voucher_type_id
             LEFT JOIN users gv ON gv.id = t.giver_id
            WHERE i.request_id = ? AND t.is_void = 0
            ORDER BY t.txn_date, t.id""",
        (req_id,),
    ).fetchall()


def decode_signature(data_url):
    """캔버스 data URL(PNG) → 검증·정규화된 PNG 바이트 (흰 배경, 최대 900x450)."""
    prefix = "data:image/png;base64,"
    if not data_url or not data_url.startswith(prefix):
        raise LedgerError("서명 이미지가 전송되지 않았습니다. 서명란에 서명해 주세요.")
    b64 = data_url[len(prefix):]
    if len(b64) > 3_000_000:
        raise LedgerError("서명 이미지가 너무 큽니다.")
    try:
        img = Image.open(BytesIO(base64.b64decode(b64, validate=True)))
        if img.format != "PNG" or img.width * img.height > 12_000_000:
            raise ValueError("bad image")
        img.load()
    except Exception as e:  # noqa: BLE001
        raise LedgerError("서명 이미지를 읽을 수 없습니다. 다시 서명해 주세요.") from e
    rgba = img.convert("RGBA")
    canvas = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
    canvas.alpha_composite(rgba)
    gray = canvas.convert("L")
    if ImageOps.invert(gray).point(lambda p: 255 if p > 60 else 0).getbbox() is None:
        raise LedgerError("서명이 비어 있습니다. 서명란에 서명해 주세요.")
    gray.thumbnail((900, 450))
    out = BytesIO()
    gray.save(out, "PNG", optimize=True)
    return out.getvalue()


def _norm_name(s):
    return re.sub(r"\s+", "", s or "")


def complete_signature(db, req_id, data_url, signer_name, user_agent):
    png = decode_signature(data_url)
    with write_tx(db):
        req = db.execute("SELECT * FROM signature_requests WHERE id = ?", (req_id,)).fetchone()
        if req is None or req["status"] != "PENDING":
            raise LedgerError("이미 처리되었거나 취소된 서명 요청입니다.")
        if req["expires_at"] < now_str():
            raise LedgerError("서명 기한이 지났습니다. 담당자에게 재요청을 부탁하세요.")
        lines = request_lines(db, req_id)
        if not lines:
            raise LedgerError("서명할 내역이 없습니다. (모두 취소됨)")
        recipient = db.execute("SELECT * FROM users WHERE id = ?", (req["recipient_id"],)).fetchone()
        if _norm_name(signer_name) != _norm_name(recipient["name"]):
            raise LedgerError("입력한 성명이 수령자 정보와 일치하지 않습니다. 본인 성명을 정확히 입력하세요.")

        ts = now()
        fname = f"sig_{req_id}_{ts:%Y%m%d%H%M%S}_{secrets.token_hex(4)}.png"
        path = os.path.join(signature_dir(), fname)
        with open(path, "wb") as fh:
            fh.write(png)
        try:
            db.execute(
                """UPDATE signature_requests
                      SET status='SIGNED', signed_at=?, signer_name=?, signer_ip=?, signer_ua=?,
                          signature_file=?, signature_sha256=?
                    WHERE id=?""",
                (now_str(ts), signer_name.strip()[:100], client_ip(), (user_agent or "")[:300],
                 fname, sha256_hex(png), req_id),
            )
            ids = [r["id"] for r in lines]
            db.execute(
                f"UPDATE transactions SET sign_status='SIGNED', sign_request_id=?, updated_at=? "
                f"WHERE id IN ({','.join('?' * len(ids))})",
                [req_id, now_str(ts), *ids],
            )
            audit(db, "SIGN_COMPLETE", f"sigreq:{req_id}", f"내역 {ids}", user_id=recipient["id"])
        except BaseException:
            os.remove(path)
            raise


# --------------------------------------------------------------------------- 엑셀 내보내기

SIGN_LABELS = {"NONE": "", "UNREQUESTED": "미요청", "REQUESTED": "서명대기", "SIGNED": "서명완료"}


def export_xlsx(rows, stock, title):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "수불대장"
    headers = [
        ("No", 6), ("일자", 11), ("구분", 6), ("상품권 종류", 16), ("권종", 10), ("수량", 7),
        ("입고금액", 12), ("출고금액", 12), ("권종재고(매)", 11), ("총잔액", 13), ("일련번호", 18),
        ("구매처", 14), ("지급자", 10), ("수령자", 10), ("부서", 12), ("수령자 이메일", 24), ("용도/비고", 24),
        ("서명상태", 9), ("서명일시", 19), ("서명", 18), ("서명파일 SHA-256", 20), ("취소", 6), ("취소사유", 18),
    ]
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    ws.append([h for h, _ in headers])
    for i, (_h, w) in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), head_fill, border
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w

    sig_col = [h for h, _ in headers].index("서명") + 1
    sig_dir = signature_dir()
    for row_idx, r in enumerate(rows, start=2):
        is_in = r["direction"] == "IN"
        ws.append([
            r["id"], r["txn_date"], "입고" if is_in else "출고", r["type_name"], r["face_value"], r["quantity"],
            r["amount"] if is_in else None, None if is_in else r["amount"], r["qty_balance"], r["amount_balance"],
            r["serial_no"], r["counterparty"], r["giver_name"], r["recipient_name"], r["recipient_dept"],
            r["recipient_email"], r["purpose"], SIGN_LABELS.get(r["sign_status"], ""),
            r["signed_at"] if r["sign_status"] == "SIGNED" else None, None,
            r["signature_sha256"] if r["sign_status"] == "SIGNED" else None,
            "취소" if r["is_void"] else "", r["void_reason"],
        ])
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col)
            c.border = border
            c.alignment = Alignment(vertical="center")
            if col in (5, 6, 7, 8, 9, 10):
                c.number_format = "#,##0"
            if r["is_void"]:
                c.font = Font(strike=True, color="999999")
        if r["sign_status"] == "SIGNED" and r["signature_file"]:
            path = os.path.join(sig_dir, r["signature_file"])
            if os.path.exists(path):
                img = XLImage(path)
                ratio = 40 / img.height
                img.height, img.width = 40, int(img.width * ratio)
                ws.add_image(img, f"{get_column_letter(sig_col)}{row_idx}")
                ws.row_dimensions[row_idx].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("재고현황")
    ws2.append(["상품권 종류", "권종", "입고수량", "출고수량", "현재고(매)", "잔액(원)"])
    for s in stock:
        ws2.append([s["name"], s["face_value"], s["in_qty"], s["out_qty"], s["qty"], s["balance"]])
    ws2.append(["합계", None, None, None, sum(s["qty"] for s in stock), sum(s["balance"] for s in stock)])
    for col, w in zip("ABCDEF", (18, 10, 10, 10, 12, 14)):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=1):
        for c in row:
            c.border = border
            if c.row == 1:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            elif c.column >= 2:
                c.number_format = "#,##0"
    wb.properties.title = title

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
