"""관리자: 직원 계정, 상품권 종류/권종, 감사 로그."""
import csv
import io
import re

from flask import Blueprint, current_app, flash, g, redirect, render_template, request, send_file, url_for

from .auth import EMAIL_RE, admin_required
from .db import get_db, write_tx
from .services import audit
from .utils import now_str, parse_int

bp = Blueprint("admin", __name__, url_prefix="/admin")


def _clean_user(form):
    email = (form.get("email") or "").strip().lower()
    name = re.sub(r"\s+", " ", form.get("name") or "").strip()
    if not EMAIL_RE.match(email):
        raise ValueError(f"메일 주소가 올바르지 않습니다: {email or '(빈 값)'}")
    domains = current_app.config["ALLOWED_EMAIL_DOMAINS"]
    if domains and email.split("@", 1)[1] not in domains:
        raise ValueError(f"허용되지 않은 메일 도메인입니다: {email}")
    if not name:
        raise ValueError(f"성명을 입력하세요: {email}")
    role = form.get("role") if form.get("role") in ("admin", "staff") else "staff"
    return dict(email=email, name=name[:50], department=(form.get("department") or "").strip()[:100],
                employee_no=(form.get("employee_no") or "").strip()[:50], role=role)


# --------------------------------------------------------------------------- 직원

@bp.get("/users")
@admin_required
def users():
    rows = get_db().execute(
        """SELECT u.*, (SELECT COUNT(*) FROM transactions t WHERE t.recipient_id = u.id AND t.is_void = 0) AS received
             FROM users u ORDER BY u.is_active DESC, u.name"""
    ).fetchall()
    return render_template("admin/users.html", users=rows)


@bp.post("/users")
@admin_required
def user_create():
    db = get_db()
    try:
        u = _clean_user(request.form)
        with write_tx(db):
            if db.execute("SELECT 1 FROM users WHERE email = ?", (u["email"],)).fetchone():
                raise ValueError(f"이미 등록된 메일 주소입니다: {u['email']}")
            ts = now_str()
            db.execute(
                """INSERT INTO users (email, name, department, employee_no, role, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (u["email"], u["name"], u["department"], u["employee_no"], u["role"], ts, ts),
            )
            audit(db, "USER_CREATE", u["email"], f"{u['name']} / {u['role']}")
        flash(f"{u['name']} ({u['email']}) 등록 완료", "success")
    except ValueError as e:
        flash(str(e), "error")
    return redirect(url_for("admin.users"))


@bp.post("/users/<int:uid>")
@admin_required
def user_update(uid):
    db = get_db()
    try:
        u = _clean_user(request.form)
        active = 1 if request.form.get("is_active") else 0
        with write_tx(db):
            dup = db.execute("SELECT id FROM users WHERE email = ? AND id != ?", (u["email"], uid)).fetchone()
            if dup:
                raise ValueError(f"이미 다른 직원이 사용 중인 메일 주소입니다: {u['email']}")
            db.execute(
                """UPDATE users SET email=?, name=?, department=?, employee_no=?, role=?, is_active=?, updated_at=?
                    WHERE id=?""",
                (u["email"], u["name"], u["department"], u["employee_no"], u["role"], active, now_str(), uid),
            )
            admins = db.execute("SELECT COUNT(*) FROM users WHERE role = 'admin' AND is_active = 1").fetchone()[0]
            if admins == 0:
                raise ValueError("활성 관리자가 최소 1명은 있어야 합니다.")
            if uid == g.user["id"] and (u["role"] != "admin" or not active):
                raise ValueError("본인의 관리자 권한/사용 상태는 직접 해제할 수 없습니다.")
            audit(db, "USER_UPDATE", u["email"], f"{u['name']} / {u['role']} / active={active}")
        flash(f"{u['name']} 정보를 수정했습니다.", "success")
    except ValueError as e:
        flash(str(e), "error")
    return redirect(url_for("admin.users"))


ROSTER_HEADERS = ["이메일", "성명", "부서", "사번", "권한"]


def _roster_rows(file):
    """업로드한 명단(.xlsx 또는 .csv)의 행 목록. 형식 오류는 ValueError."""
    raw = file.read()
    if file.filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ValueError("엑셀 파일을 읽을 수 없습니다.") from e
        ws = wb.worksheets[0]
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(max_row=5000, values_only=True)]
        wb.close()
        return rows
    for enc in ("utf-8-sig", "cp949"):
        try:
            return list(csv.reader(io.StringIO(raw.decode(enc))))
        except UnicodeDecodeError:
            continue
    raise ValueError("파일 인코딩을 읽을 수 없습니다. 엑셀(.xlsx) 또는 CSV 로 저장하세요.")


@bp.get("/users/export.xlsx")
@admin_required
def user_export():
    """현재 직원 명단을 업로드 양식 그대로 내려받기 (고쳐서 다시 올리면 반영)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "직원명단"
    ws.append(ROSTER_HEADERS + ["사용", "Microsoft 계정 연결"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for u in get_db().execute("SELECT * FROM users ORDER BY is_active DESC, name").fetchall():
        ws.append([u["email"], u["name"], u["department"], u["employee_no"],
                   "관리자" if u["role"] == "admin" else "", "" if u["is_active"] else "중지",
                   "연결됨" if u["sso_subject"] else ""])
    for col, w in zip("ABCDEFG", (28, 12, 16, 10, 8, 6, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"직원명단_{now_str()[:10].replace('-', '')}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/users/import")
@admin_required
def user_import():
    """명단 일괄 등록: 이메일, 성명, 부서, 사번[, 권한] — 엑셀(.xlsx) 또는 CSV. 없는 직원은 지우지 않습니다."""
    file = request.files.get("file")
    if not file or not file.filename:
        flash("엑셀(.xlsx) 또는 CSV 파일을 선택하세요.", "error")
        return redirect(url_for("admin.users"))
    try:
        rows = _roster_rows(file)
    except ValueError as e:
        flash(str(e), "error")
        return redirect(url_for("admin.users"))

    db = get_db()
    added = updated = 0
    errors = []
    ts = now_str()
    with write_tx(db):
        for lineno, row in enumerate(rows, start=1):
            if not row or not any(c.strip() for c in row) or "@" not in row[0]:
                continue  # 빈 줄 / 머리글
            cols = (list(row) + [""] * 5)[:5]
            role_raw = cols[4].strip().lower()
            form = dict(email=cols[0], name=cols[1], department=cols[2], employee_no=cols[3],
                        role="admin" if role_raw in ("admin", "관리자") else "staff")
            try:
                u = _clean_user(form)
            except ValueError as e:
                errors.append(f"{lineno}행: {e}")
                continue
            exists = db.execute("SELECT id FROM users WHERE email = ?", (u["email"],)).fetchone()
            if exists:  # 기존 직원은 이름/부서/사번만 갱신 (권한·사용 여부는 화면에서 직접 변경)
                db.execute("UPDATE users SET name=?, department=?, employee_no=?, updated_at=? WHERE id=?",
                           (u["name"], u["department"], u["employee_no"], ts, exists["id"]))
                updated += 1
            else:
                db.execute(
                    """INSERT INTO users (email, name, department, employee_no, role, created_at, updated_at)
                       VALUES (?, ?, ?, ?, ?, ?, ?)""",
                    (u["email"], u["name"], u["department"], u["employee_no"], u["role"], ts, ts),
                )
                added += 1
        audit(db, "USER_IMPORT", file.filename, f"추가 {added}, 갱신 {updated}, 오류 {len(errors)}")
    flash(f"명단 등록 결과: 신규 {added}명, 갱신 {updated}명", "success")
    for e in errors[:20]:
        flash(e, "error")
    return redirect(url_for("admin.users"))


# --------------------------------------------------------------------------- 상품권 종류/권종

@bp.get("/types")
@admin_required
def types():
    rows = get_db().execute(
        """SELECT vt.*, (SELECT COUNT(*) FROM transactions t WHERE t.voucher_type_id = vt.id) AS used
             FROM voucher_types vt ORDER BY vt.sort_order, vt.name, vt.face_value DESC"""
    ).fetchall()
    return render_template("admin/types.html", types=rows)


@bp.post("/types")
@admin_required
def type_create():
    db = get_db()
    name = (request.form.get("name") or "").strip()[:100]
    face = parse_int(request.form.get("face_value"))
    if not name or not face or face <= 0:
        flash("상품권 종류와 권종 금액을 입력하세요.", "error")
        return redirect(url_for("admin.types"))
    with write_tx(db):
        if db.execute("SELECT 1 FROM voucher_types WHERE name = ? AND face_value = ?", (name, face)).fetchone():
            flash(f"이미 등록된 권종입니다: {name} {face:,}원", "error")
            return redirect(url_for("admin.types"))
        db.execute("INSERT INTO voucher_types (name, face_value, sort_order, created_at) VALUES (?, ?, ?, ?)",
                   (name, face, parse_int(request.form.get("sort_order"), 0), now_str()))
        audit(db, "TYPE_CREATE", name, f"{face:,}원")
    flash(f"{name} {face:,}원권을 추가했습니다.", "success")
    return redirect(url_for("admin.types"))


@bp.post("/types/<int:type_id>")
@admin_required
def type_update(type_id):
    """종류명/정렬/사용여부만 변경 가능. 권종 금액은 기존 기록 보호를 위해 수정하지 않습니다."""
    db = get_db()
    name = (request.form.get("name") or "").strip()[:100]
    if not name:
        flash("상품권 종류명을 입력하세요.", "error")
        return redirect(url_for("admin.types"))
    with write_tx(db):
        cur = db.execute("SELECT face_value FROM voucher_types WHERE id = ?", (type_id,)).fetchone()
        dup = cur and db.execute("SELECT 1 FROM voucher_types WHERE name = ? AND face_value = ? AND id != ?",
                                 (name, cur["face_value"], type_id)).fetchone()
        if cur is None or dup:
            flash("같은 이름·권종이 이미 있거나 존재하지 않는 항목입니다.", "error")
            return redirect(url_for("admin.types"))
        db.execute("UPDATE voucher_types SET name=?, sort_order=?, is_active=? WHERE id=?",
                   (name, parse_int(request.form.get("sort_order"), 0),
                    1 if request.form.get("is_active") else 0, type_id))
        audit(db, "TYPE_UPDATE", f"type:{type_id}", name)
    flash("수정했습니다.", "success")
    return redirect(url_for("admin.types"))


# --------------------------------------------------------------------------- 감사 로그

@bp.get("/audit")
@admin_required
def audit_log():
    rows = get_db().execute(
        """SELECT a.*, u.name AS user_name, u.email AS user_email
             FROM audit_log a LEFT JOIN users u ON u.id = a.user_id
            ORDER BY a.id DESC LIMIT 1000"""
    ).fetchall()
    return render_template("admin/audit.html", rows=rows)
